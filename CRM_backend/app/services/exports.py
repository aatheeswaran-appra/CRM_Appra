"""Generate reports in memory; never persist report data or export files."""

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.enums import ExportFormat
from app.schemas import ReportOut

Cell = str | int | float | None
Row = list[Cell]


def safe_text(value: Cell) -> Cell:
    # CSV formula injection protection, including leading whitespace.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def sections(report: ReportOut) -> dict[str, list[Row]]:
    return {
        "Period": [["Field", "Value"]]
        + [[key, str(value)] for key, value in report.period.model_dump(by_alias=True).items()],
        "Summary": [["Metric", "Value"]]
        + [[key, value] for key, value in report.summary.model_dump(by_alias=True).items()],
        "Lead Growth": [["Label", "Count"]]
        + [[item.label, item.count] for item in report.lead_growth],
        "Lead Sources": [["Source", "Count", "Percentage"]]
        + [[item.source.value, item.count, item.percentage] for item in report.lead_sources],
        "Status Breakdown": [["Status", "Count"]]
        + [[item.status.value, item.count] for item in report.status_breakdown],
        "Top Requirements": [["Requirement", "Count"]]
        + [[item.requirement, item.count] for item in report.top_requirements],
        "Follow-up Performance": [["Metric", "Value"]]
        + [
            [key, value]
            for key, value in report.follow_up_performance.model_dump(by_alias=True).items()
        ],
        "Highlights": [["Metric", "Message", "Current", "Previous", "Change", "Unit"]]
        + [
            [item.metric, item.message, item.current, item.previous, item.change, item.unit]
            for item in report.highlights
        ],
    }


def export_report(report: ReportOut, export_format: ExportFormat) -> tuple[bytes, str, str]:
    filename = f"appra-report-{report.period.from_date}-{report.period.to}.{export_format.value}"
    data = sections(report)
    if export_format == ExportFormat.CSV:
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(
            [
                "Section",
                "Field / Metric",
                "Value / Count",
                "Percentage / Current",
                "Previous",
                "Change",
                "Unit",
            ]
        )
        for section, rows in data.items():
            for row in rows[1:]:
                cells = [safe_text(cell) for cell in row]
                writer.writerow([section, *cells, *[""] * (6 - len(cells))])
        return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", filename
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for section, rows in data.items():
        sheet = workbook.create_sheet(section)
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # User-entered requirements must never become formulas.
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 65 if section == "Highlights" else 30
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return (
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename,
    )
