import csv
from datetime import datetime
from io import BytesIO, StringIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, func, inspect, select
from sqlalchemy.orm import Session

from app.enums import FollowUpStatus, LeadStatus
from app.models import Customer, FollowUp
from app.seed import seed
from tests.conftest import NOW

pytestmark = pytest.mark.integration


def ist(value: str) -> datetime:
    return datetime.fromisoformat(value + "+05:30")


def add_report_data(engine: Engine) -> None:
    with Session(engine) as session, session.begin():
        definitions = [
            ("Arun", "2026-09-01T00:00", "NEW", "WEBSITE", "Website Development", None),
            ("Priya", "2026-09-08T00:00", "INTERESTED", "WHATSAPP", "Digital Marketing", None),
            (
                "Rajesh",
                "2026-09-15T10:00",
                "CLOSED",
                "WEBSITE",
                "Website Development",
                "2026-09-16T10:00",
            ),
            ("Abi", "2026-09-18T11:00", "INTERESTED", "REFERRAL", "SEO", None),
            (
                "Old converted",
                "2026-08-02T10:00",
                "CLOSED",
                "CALL",
                "Website Development",
                "2026-09-10T10:00",
            ),
            ("August", "2026-08-31T23:59", "NEW", "OTHER", "SEO", None),
            ("July", "2026-07-10T10:00", "CLOSED", "OTHER", "SEO", "2026-08-05T10:00"),
        ]
        for index, (name, created_at, status, source, requirement, closed_at) in enumerate(
            definitions
        ):
            session.add(
                Customer(
                    name=name,
                    phone=f"90000000{index:02}",
                    requirement=requirement,
                    source=source,
                    status=LeadStatus(status),
                    closed_at=ist(closed_at) if closed_at else None,
                    created_at=ist(created_at),
                )
            )
        session.flush()
        for scheduled, completed, reschedules in [
            ("2026-09-18T10:00", None, 0),
            ("2026-09-01T10:00", "2026-09-01T12:00", 0),
            ("2026-09-08T10:00", None, 2),
            ("2026-09-20T10:00", None, 1),
            (
                "2026-09-25T10:00",
                "2026-09-18T11:00",
                0,
            ),  # Future due: exclude from completion denominator.
            ("2026-09-18T12:00", "2026-09-18T12:00", 0),
            ("2026-08-05T10:00", "2026-08-05T12:00", 0),
            ("2026-08-08T10:00", None, 0),
            ("2026-08-09T10:00", None, 0),
            ("2026-08-10T10:00", None, 0),
        ]:
            session.add(
                FollowUp(
                    customer_id=1,
                    follow_up_at=ist(scheduled),
                    preferred_contact="CALL",
                    status=FollowUpStatus.COMPLETED if completed else FollowUpStatus.PENDING,
                    completed_at=ist(completed) if completed else None,
                    reschedule_count=reschedules,
                )
            )


def test_dashboard_counts_and_closed_at_comparisons(client: TestClient, engine: Engine) -> None:
    add_report_data(engine)
    response = client.get("/api/dashboard")
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["stats"] == {
        "totalCustomers": 7,
        "totalCustomersChange": 100,
        "followUpsToday": 1,
        "overdueFollowUps": 5,
        "closedCustomers": 3,
        "closedCustomersChange": 100,
    }
    assert len(data["recentCustomers"]) == 5
    assert data["recentCustomers"][0]["name"] == "Abi"
    assert len(data["overdueFollowUps"]) == 5
    assert len(data["todayFollowUps"]) == 1


def test_reports_exact_cohort_math_growth_and_highlights(
    client: TestClient, engine: Engine
) -> None:
    add_report_data(engine)
    response = client.get("/api/reports?from=2026-09-01&to=2026-09-30")
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["period"] == {
        "from": "2026-09-01",
        "to": "2026-09-30",
        "timezone": "Asia/Kolkata",
        "comparisonFrom": "2026-08-01",
        "comparisonTo": "2026-08-31",
    }
    assert data["summary"] == {
        "totalLeads": 4,
        "totalLeadsChange": 100,
        "interestedLeads": 2,
        "interestedLeadsChange": 100,
        "conversionRate": 25,
        "conversionRateChange": 25,
        "followUpCompletionRate": 50,
        "followUpCompletionRateChange": 25,
    }
    assert [item["count"] for item in data["leadGrowth"]] == [1, 1, 2, 0, 0]
    assert next(item for item in data["leadSources"] if item["source"] == "WEBSITE") == {
        "source": "WEBSITE",
        "count": 2,
        "percentage": 50,
    }
    assert sum(item["count"] for item in data["statusBreakdown"]) == 4
    assert data["topRequirements"][0] == {"requirement": "Website Development", "count": 2}
    assert data["followUpPerformance"] == {
        "completedToday": 2,
        "upcoming": 1,
        "overdue": 5,
        "rescheduled": 2,
        "completionRate": 50,
    }
    conversion = next(item for item in data["highlights"] if item["metric"] == "conversionRate")
    assert conversion["change"] == 25 and conversion["unit"] == "percentage_points"
    assert client.get("/api/reports").json()["data"] == data


def test_custom_range_uses_equal_length_comparison(client: TestClient, engine: Engine) -> None:
    add_report_data(engine)
    response = client.get("/api/reports?from=2026-09-08&to=2026-09-18")
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["period"]["comparisonFrom"] == "2026-08-28"
    assert data["period"]["comparisonTo"] == "2026-09-07"
    assert data["summary"]["totalLeads"] == 3
    assert len(data["leadGrowth"]) == 2


def test_empty_reports_and_dashboard(client: TestClient) -> None:
    data = client.get("/api/reports").json()["data"]
    assert all(value == 0 for value in data["summary"].values())
    assert data["highlights"] == [] and data["topRequirements"] == []
    assert len(data["leadSources"]) == 6 and len(data["statusBreakdown"]) == 6
    assert all(item["count"] == 0 for item in data["leadGrowth"])
    assert all(
        value == 0 for value in client.get("/api/dashboard").json()["data"]["stats"].values()
    )


def test_exports_include_report_sections_and_preserve_values(
    client: TestClient, engine: Engine
) -> None:
    add_report_data(engine)
    csv_response = client.get("/api/reports/export?format=csv")
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert "appra-report-2026-09-01-2026-09-30.csv" in csv_response.headers["content-disposition"]
    rows = list(csv.reader(StringIO(csv_response.content.decode("utf-8-sig"))))
    assert ["Summary", "totalLeads", "4", "", "", "", ""] in rows
    assert {
        "Period",
        "Summary",
        "Lead Growth",
        "Lead Sources",
        "Status Breakdown",
        "Top Requirements",
        "Follow-up Performance",
        "Highlights",
    } <= {row[0] for row in rows}
    response = client.get("/api/reports/export?format=xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert len(workbook.sheetnames) == 8
    assert workbook["Summary"]["B2"].value == 4
    assert workbook["Lead Sources"]["B2"].value == 2
    workbook.close()


def test_exports_do_not_execute_customer_requirement_formulas(
    client: TestClient, payload: dict[str, object]
) -> None:
    assert (
        client.post(
            "/api/customers",
            json={**payload, "requirement": '=HYPERLINK("https://example.invalid","test")'},
        ).status_code
        == 201
    )
    csv_response = client.get("/api/reports/export")
    assert "'=HYPERLINK" in csv_response.content.decode("utf-8-sig")
    response = client.get("/api/reports/export?format=xlsx")
    workbook = load_workbook(BytesIO(response.content))
    cell = workbook["Top Requirements"]["A2"]
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"
    workbook.close()


def test_seed_is_repeatable_and_only_two_business_tables_exist(engine: Engine) -> None:
    with Session(engine) as session:
        assert seed(session, NOW) == 32
        assert seed(session, NOW) == 0
        assert session.scalar(select(func.count(Customer.id))) == 32
        assert session.scalar(select(func.count(FollowUp.id))) == 32
    assert set(inspect(engine).get_table_names()) == {"customers", "follow_ups", "alembic_version"}


def test_migration_matches_models(engine: Engine) -> None:
    from alembic.autogenerate import compare_metadata
    from alembic.migration import MigrationContext

    from app.models import Base

    with engine.connect() as connection:
        context = MigrationContext.configure(connection, opts={"compare_type": True})
        assert compare_metadata(context, Base.metadata) == []
